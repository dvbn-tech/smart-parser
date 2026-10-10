from pathlib import Path

import openpyxl
import pytest

from parser.importers.customer_xlsx import derive_masks, read_workbook, username_from


def link(ws, coord: str, text: str, url: str) -> None:  # type: ignore[no-untyped-def]
    ws[coord] = text
    ws[coord].hyperlink = url


@pytest.fixture
def workbook(tmp_path: Path) -> Path:
    wb = openpyxl.Workbook()
    comp = wb.active
    comp.title = "конкруенты"
    comp["A1"], comp["B1"], comp["C1"] = "название бренда", "MELA", "laish.brand\n"
    comp["A2"] = "аккаунт бренда"
    comp["B2"] = "https://www.instagram.com/sport.mela/#"
    link(
        comp,
        "C2",
        "https://www.instagram.com/laish.brand?igsh=X#",
        "https://www.instagram.com/laish.brand?igsh=X",
    )
    comp["A3"] = "креаторы"
    for i, u in enumerate(
        ["katrin_mela", "aani.mela/reels/", "arina.mela", "katrin_mela"], start=3
    ):
        comp[f"B{i}"] = f"https://www.instagram.com/{u}"
    comp["A8"] = "ключевые слова"
    comp["A9"] = "название бренда"  # подписи ниже по листу не должны сбивать разбор
    own = wb.create_sheet("наши креаторы")
    own["A1"], own["B1"] = "аккаунт бренда", "https://www.instagram.com/15.wear?igsh=abc"
    own["A2"], own["B2"], own["C2"] = "ФИ креатора", "ссылка на аккаунт", "куратор"
    link(
        own, "B3", "lerok.onefivе", "https://www.instagram.com/lerok.onefive?stkn=1"
    )  # русская «е» в тексте
    own["A3"], own["C3"] = "Сирецкая Валерия", "Куратор Оля"
    link(own, "B4", "alina", "https://www.instagram.com/reel/DeEv3FaRtzt/")
    own["A4"] = "Пивоварова Алина"
    link(own, "B5", "kris.prosport", "https://www.instagram.com/kris_wears?stkn=2")
    own["A5"], own["C5"] = "Ковалева Эллина (обзор)", "Куратор Женя"
    kw = wb.create_sheet("ключевые слова")
    kw["E13"] = "у креаторов бренда MELA окончание названия бренда - mela\nпример: valerchik_mela"
    stats = wb.create_sheet("Статистика февраль")
    stats["B1"], stats["E1"] = "Product ID", "Substitution Article"
    stats["B3"], stats["E3"] = 152154933, "WW261170"
    wb.create_sheet("март Макс")["A1"] = "статус"
    path = tmp_path / "lists.xlsx"
    wb.save(path)
    return path


def test_read_workbook(workbook: Path) -> None:
    plan = read_workbook(workbook)
    brands = {b.name: b for b in plan.brands}
    assert set(brands) == {"MELA", "laish.brand", "15.wear"}

    mela = brands["MELA"]
    assert mela.shop is not None and mela.shop.username == "sport.mela"
    assert [c.username for c in mela.creators] == [
        "katrin_mela",
        "aani.mela",
        "arina.mela",
    ]  # дубль убран
    assert "*[._]mela*" in mela.masks

    own = brands["15.wear"]
    assert own.is_own and own.ww_codes == {"WW261170"} and own.wb_articles == {"152154933"}
    creators = {c.username: c for c in own.creators}
    assert set(creators) == {"lerok.onefive", "kris_wears"}  # ссылка на ролик не превращается в ник
    assert creators["lerok.onefive"].curator == "Оля"
    assert creators["kris_wears"].link_type == "reviewer"

    messages = " ".join(i.message for i in plan.issues)
    assert "reel/DeEv3FaRtzt" in messages
    assert "kris.prosport" in messages
    assert "laish.brand" in " ".join(i.where for i in plan.issues)  # бренд без креаторов
    assert plan.skipped_sheets == ["март Макс"]


@pytest.mark.parametrize(
    ("text", "url", "expected"),
    [
        ("x", "https://www.instagram.com/Name.One/reels/?igsh=1", "name.one"),
        ("ann_onefive рабочий", None, "ann_onefive"),
        ("alina", "https://www.instagram.com/reel/ABC/", None),
        ("lerok.onefivе", None, "lerok.onefive"),
        ("https://www.instagram.com/sport.mela/#", None, "sport.mela"),
    ],
)
def test_username_from(text: str, url: str | None, expected: str | None) -> None:
    assert username_from(text, url)[0] == expected


def test_derive_masks() -> None:
    assert derive_masks(["a.base", "b_base", "c.base", "d.other"]) == ["*[._]base*"]
    assert derive_masks(["svoya.veronika", "svoya.ann"]) == ["*svoya*"]
    assert derive_masks(["llena.wb", "kiraa.wb", "nastiy.wb"]) == []  # «wb» — общее слово
