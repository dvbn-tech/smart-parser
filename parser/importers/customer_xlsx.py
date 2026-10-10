"""Импорт списков заказчика из его таблицы «дашборд, анализ конкурентов.xlsx».

Что берём:
  • «конкуренты» — бренды (строка «название бренда»), их аккаунты («аккаунт бренда»)
    и известные креаторы (столбец под брендом, начиная со строки «креаторы»);
  • «наши креаторы» — наш бренд (B1) и наши креаторы: ФИ, ник/ссылка, куратор;
  • «ключевые слова» — правила заказчика: окончания ников креаторов бренда и сайт бренда;
  • листы статистики WB и закупок — подменные WW-коды и артикулы WB нашего бренда.

Чтение — чистая функция read_workbook() → ImportPlan с отчётом о проблемах в данных;
запись в базу — apply_plan(). Ник берётся из гиперссылки ячейки, а не из её текста:
текст часто расходится со ссылкой (переименование, опечатки).
"""

import re
from collections import Counter
from collections.abc import Iterable
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any
from urllib.parse import urlparse

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

USERNAME_RE = re.compile(r"^[a-z0-9._]{1,30}$")
RESERVED_PATHS = {"reel", "reels", "p", "tv", "stories", "explore", "accounts", "direct"}
# Кириллические буквы, неотличимые на глаз от латинских (lerok.onefivе с русской «е»)
HOMOGLYPHS = str.maketrans("аеорсухкмтвнАЕОРСУХКМТВН", "aeopcyxkmtbhAEOPCYXKMTBH")
GENERIC_SEGMENTS = {
    "wb",
    "by",
    "the",
    "official",
    "shop",
    "style",
    "reels",
    "sport",
    "fit",
    "obzor",
}
WW_RE = re.compile(r"^WW\d{5,9}$", re.IGNORECASE)


@dataclass
class Issue:
    where: str
    message: str
    level: str = "warn"  # warn — нужно проверить; info — к сведению


@dataclass
class AccountRef:
    username: str
    role: str  # shop | creator
    source: str
    person_name: str | None = None
    curator: str | None = None
    link_type: str | None = None


@dataclass
class BrandPlan:
    name: str
    is_own: bool
    shop: AccountRef | None = None
    creators: list[AccountRef] = field(default_factory=list)
    masks: list[str] = field(default_factory=list)
    domains: list[str] = field(default_factory=list)
    ww_codes: set[str] = field(default_factory=set)
    wb_articles: set[str] = field(default_factory=set)


@dataclass
class ImportPlan:
    brands: list[BrandPlan]
    issues: list[Issue]
    skipped_sheets: list[str]


def username_from(text: Any, link: str | None) -> tuple[str | None, list[str]]:
    """Ник из ссылки (приоритет) или текста ячейки. Возвращает (ник, заметки)."""
    notes: list[str] = []
    from_link = _parse_profile_url(link) if link else None
    if link and from_link is None and "instagram.com" in link:
        # ссылка на ролик или пост: текст ячейки — обычно имя, а не ник; не угадываем
        notes.append(f"ссылка ведёт не на профиль, а на {link.split('?')[0]}")
        return None, notes
    candidates = [c for c in (_parse_profile_url(t) for t in _text_tokens(text)) if c]
    if from_link:
        others = [c for c in candidates if c != from_link]
        if others:
            notes.append(
                f"в ячейке написано «{', '.join(others)}», ссылка ведёт на @{from_link} — взяли ссылку"
            )
        return from_link, notes
    if len(candidates) > 1:
        notes.append(f"в ячейке несколько ников: {', '.join(candidates)} — взяли первый")
    return (candidates[0] if candidates else None), notes


def _text_tokens(text: Any) -> list[str]:
    if text is None:
        return []
    parts = re.split(r"[\s,;]+", str(text).strip())
    return [p for p in parts if p and p.lower() not in {"рабочий", "обзор"}]


def _parse_profile_url(value: str | None) -> str | None:
    if not value:
        return None
    v = value.strip()
    if "instagram.com" in v:
        path = urlparse(v if "://" in v else f"https://{v}").path
        parts = [p for p in path.split("/") if p]
        if not parts or parts[0].lower() in RESERVED_PATHS:
            return None
        v = parts[0]
    elif "/" in v:
        return None
    v = v.split("?")[0].split("#")[0].lstrip("@").strip().lower().translate(HOMOGLYPHS)
    return v if USERNAME_RE.match(v) else None


def _cell_link(cell: Any) -> str | None:
    link = cell.hyperlink.target if cell.hyperlink else None
    if link:
        return str(link)
    if isinstance(cell.value, str) and "instagram.com" in cell.value:
        return cell.value.strip()
    return None


def _first_col_value(ws: Worksheet, row: int) -> str:
    v = ws.cell(row=row, column=1).value
    return str(v).strip().lower() if v is not None else ""


# --- листы ----------------------------------------------------------------------------


def _read_competitors(ws: Worksheet, issues: list[Issue]) -> list[BrandPlan]:
    name_row = acc_row = creators_row = None
    # Берём первое вхождение: ниже на листе есть блок «ключевые слова» с теми же подписями
    for r in range(1, min(ws.max_row, 60) + 1):
        a = _first_col_value(ws, r)
        if a.startswith("название бренда") and name_row is None:
            name_row = r
        elif a.startswith("аккаунт бренда") and acc_row is None:
            acc_row = r
        elif a.startswith("креатор") and creators_row is None:
            creators_row = r
    if name_row is None or acc_row is None:
        issues.append(Issue(ws.title, "не нашли строки «название бренда» и «аккаунт бренда»"))
        return []
    brands: list[BrandPlan] = []
    for col in range(2, ws.max_column + 1):
        raw_name = ws.cell(row=name_row, column=col).value
        if raw_name is None or not str(raw_name).strip():
            continue
        name = str(raw_name).strip()
        acc_cell = ws.cell(row=acc_row, column=col)
        shop, notes = username_from(acc_cell.value, _cell_link(acc_cell))
        brand = BrandPlan(name=name, is_own=False)
        where = f"{ws.title}!{acc_cell.coordinate}"
        issues.extend(Issue(where, n) for n in notes)
        if shop:
            brand.shop = AccountRef(shop, "shop", where)
        else:
            issues.append(Issue(where, f"у бренда {name} нет ссылки на аккаунт"))
        if creators_row is not None:
            r = creators_row
            while r <= ws.max_row:
                a = _first_col_value(ws, r)
                if r > creators_row and a and not a.startswith("креатор"):
                    break  # дальше начинается другой блок («ключевые слова»)
                cell = ws.cell(row=r, column=col)
                if cell.value is not None or cell.hyperlink:
                    username, notes = username_from(cell.value, _cell_link(cell))
                    cwhere = f"{ws.title}!{cell.coordinate}"
                    issues.extend(Issue(cwhere, n) for n in notes)
                    if username:
                        brand.creators.append(AccountRef(username, "creator", cwhere))
                    else:
                        issues.append(
                            Issue(
                                cwhere, f"не удалось взять ник из «{str(cell.value).strip()[:60]}»"
                            )
                        )
                r += 1
        brands.append(brand)
    return brands


def _read_own(ws: Worksheet, issues: list[Issue]) -> BrandPlan | None:
    shop_cell = None
    header_row = None
    for r in range(1, min(ws.max_row, 10) + 1):
        a = _first_col_value(ws, r)
        if a.startswith("аккаунт бренда"):
            shop_cell = ws.cell(row=r, column=2)
        elif a.startswith("фи") or "креатор" in a:
            header_row = r
    if shop_cell is None:
        issues.append(Issue(ws.title, "не нашли строку «аккаунт бренда»"))
        return None
    shop, notes = username_from(shop_cell.value, _cell_link(shop_cell))
    issues.extend(Issue(f"{ws.title}!{shop_cell.coordinate}", n) for n in notes)
    if not shop:
        issues.append(
            Issue(f"{ws.title}!{shop_cell.coordinate}", "нет ссылки на аккаунт нашего бренда")
        )
        return None
    brand = BrandPlan(
        name=shop, is_own=True, shop=AccountRef(shop, "shop", f"{ws.title}!{shop_cell.coordinate}")
    )
    for r in range((header_row or 2) + 1, ws.max_row + 1):
        name_v = ws.cell(row=r, column=1).value
        cell = ws.cell(row=r, column=2)
        curator_v = ws.cell(row=r, column=3).value
        if name_v is None and cell.value is None:
            continue
        where = f"{ws.title}!{cell.coordinate}"
        username, notes = username_from(cell.value, _cell_link(cell))
        issues.extend(Issue(where, n) for n in notes)
        person = str(name_v).strip() if name_v else None
        if not username:
            issues.append(
                Issue(where, f"{person or 'креатор'}: нет ссылки на профиль — добавьте ник вручную")
            )
            continue
        lt = "reviewer" if person and "обзор" in person.lower() else None
        if person and person.lower() == "паблики":
            lt, person = "public", None
        brand.creators.append(
            AccountRef(
                username,
                "creator",
                where,
                person_name=person,
                curator=str(curator_v).replace("Куратор", "").strip() if curator_v else None,
                link_type=lt,
            )
        )
    return brand


_RULE_RE = re.compile(
    r"у креаторов бренда\s+(?P<brand>.+?)\s+(?:на\s+)?окончани[еяй]\s+(?:названия бренда|профиля)?"
    r"\s*[-–—:]\s*(?P<tokens>[^\n]+)",
    re.IGNORECASE,
)


def _read_keywords(ws: Worksheet) -> tuple[dict[str, list[str]], list[str]]:
    """Правила «у креаторов бренда X окончание - Y» и ссылки на сайты брендов."""
    rules: dict[str, list[str]] = {}
    urls: list[str] = []
    for row in ws.iter_rows():
        for c in row:
            if not isinstance(c.value, str):
                continue
            for m in _RULE_RE.finditer(c.value):
                tokens = re.split(r"\s+(?:либо|или)\s+|,", m.group("tokens"))
                clean = [t.strip().strip("-–— ").lower() for t in tokens]
                rules[m.group("brand").strip().lower()] = [
                    t for t in clean if re.fullmatch(r"[a-z0-9._]+", t)
                ]
            urls += re.findall(r"https?://[^\s)\"']+", c.value)
    return rules, urls


def _read_articles(ws: Worksheet) -> tuple[set[str], set[str]]:
    """WW-коды и артикулы WB из листов статистики и закупок (по заголовкам столбцов)."""
    ww: set[str] = set()
    wb: set[str] = set()
    ww_cols: list[int] = []
    wb_cols: list[int] = []
    for r in range(1, 4):
        for c in range(1, ws.max_column + 1):
            h = str(ws.cell(row=r, column=c).value or "").strip().lower()
            if h in {"substitution article", "подменный арт", "подменный артикул"}:
                ww_cols.append(c)
            elif h in {"product id", "арт вб", "артикул"}:
                wb_cols.append(c)
    for row in ws.iter_rows(min_row=2):
        for c in ww_cols:
            v = row[c - 1].value if c - 1 < len(row) else None
            if isinstance(v, str) and WW_RE.match(v.strip()):
                ww.add(v.strip().upper())
        for c in wb_cols:
            v = row[c - 1].value if c - 1 < len(row) else None
            if isinstance(v, int | float) and 10_000 <= v < 10**12:
                wb.add(str(int(v)))
    return ww, wb


# --- маски ников ----------------------------------------------------------------------


def _segments(username: str) -> list[str]:
    return [s for s in re.split(r"[._]+", username) if s]


def token_mask(token: str) -> str:
    """Длинный характерный сегмент (≥ 5 букв) ищем где угодно в нике, короткий — целым сегментом."""
    return f"*{token}*" if len(token) >= 5 else f"*[._]{token}*"


def derive_masks(usernames: Iterable[str], exclude: Iterable[str] = ()) -> list[str]:
    """Сегмент ника, общий для ≥ 3 креаторов бренда (или для всех, если их 2–3) → маска."""
    names = list(dict.fromkeys(usernames))
    excluded = {e.lower() for e in exclude}
    counts = Counter(seg for u in names for seg in set(_segments(u)))
    masks = []
    for seg, n in counts.most_common():
        enough = n >= 3 or (n >= 2 and n == len(names))
        if (
            not enough
            or len(seg) < 3
            or seg in GENERIC_SEGMENTS
            or seg in excluded
            or seg.isdigit()
        ):
            continue
        # сегмент в начале ника у всех («svoya.veronika», «svoya.ann») — маска по началу
        if all(_segments(u)[0] == seg for u in names if seg in _segments(u)) and len(seg) < 5:
            masks.append(f"{seg}[._]*")
        else:
            masks.append(token_mask(seg))
    return masks


def customer_mask(token: str) -> str:
    """Короткое окончание (≤ 4 букв) — только целым сегментом, иначе слишком много ложных."""
    return f"*[._]{token}" if len(token) <= 4 else f"*{token}"


def _matches(mask: str, usernames: Iterable[str]) -> int:
    import fnmatch

    return sum(fnmatch.fnmatchcase(u, mask) for u in usernames)


# --- сборка ---------------------------------------------------------------------------


def read_workbook(path: Path) -> ImportPlan:
    wb = openpyxl.load_workbook(path, data_only=True)
    issues: list[Issue] = []
    brands: list[BrandPlan] = []
    own: BrandPlan | None = None
    rules: dict[str, list[str]] = {}
    urls: list[str] = []
    ww: set[str] = set()
    wb_articles: set[str] = set()
    used: set[str] = set()

    for ws in wb.worksheets:
        title = ws.title.lower()
        if "конкур" in title or "конкру" in title:
            brands = _read_competitors(ws, issues)
            used.add(ws.title)
        elif "наши креатор" in title:
            own = _read_own(ws, issues)
            used.add(ws.title)
        elif "ключев" in title:
            rules, urls = _read_keywords(ws)
            used.add(ws.title)
    for ws in wb.worksheets:
        if ws.title in used:
            continue
        w, a = _read_articles(ws)
        if w or a:
            ww |= w
            wb_articles |= a
            used.add(ws.title)
    skipped = [ws.title for ws in wb.worksheets if ws.title not in used]

    if own is not None:
        own.ww_codes, own.wb_articles = ww, wb_articles
        brands.append(own)
        if ww or wb_articles:
            issues.append(
                Issue(
                    "статистика WB",
                    f"{len(ww)} подменных WW-кодов и {len(wb_articles)} артикулов WB отнесены к нашему бренду "
                    f"{own.name}: ролик с ними считается роликом бренда",
                    "info",
                )
            )

    for b in brands:
        if not b.is_own and own is not None and "munique" in b.name.lower():
            issues.append(
                Issue(
                    b.name,
                    f"в ТЗ «наш бренд» — Munique Base, а в таблице MUNIQUE BASE на листе конкурентов, "
                    f"а «наши креаторы» — у {own.name}. Импортируем {own.name} как наш, MUNIQUE BASE — как "
                    "конкурента; если оба наши — отметьте в настройках бренда",
                )
            )
    _dedupe(brands, issues)
    _apply_rules(brands, rules, urls, issues)
    for b in brands:
        if not b.creators:
            issues.append(
                Issue(
                    b.name, "в списке нет ни одного креатора — их найдёт только автопоиск", "info"
                )
            )
    return ImportPlan(brands=brands, issues=issues, skipped_sheets=skipped)


def _dedupe(brands: list[BrandPlan], issues: list[Issue]) -> None:
    owners: dict[str, str] = {}
    for b in brands:
        seen: dict[str, AccountRef] = {}
        for c in b.creators:
            if c.username in seen:
                issues.append(
                    Issue(
                        c.source,
                        f"@{c.username} уже есть в списке {b.name} ({seen[c.username].source})",
                        "info",
                    )
                )
                continue
            seen[c.username] = c
            if c.username in owners and owners[c.username] != b.name:
                issues.append(
                    Issue(c.source, f"@{c.username} указан и у {owners[c.username]}, и у {b.name}")
                )
            owners.setdefault(c.username, b.name)
        b.creators = list(seen.values())


def _apply_rules(
    brands: list[BrandPlan], rules: dict[str, list[str]], urls: list[str], issues: list[Issue]
) -> None:
    def key(name: str) -> str:
        return re.sub(r"[\W_]+", "", name.lower())

    for b in brands:
        creators = [c.username for c in b.creators]
        shop_segments = set(_segments(b.shop.username)) if b.shop else set()
        rule = next(
            (
                t
                for name, t in rules.items()
                if key(name) == key(b.name) or key(name) in key(b.name)
            ),
            [],
        )
        for token in rule:
            mask = customer_mask(token)
            hits = _matches(mask, creators)
            if hits:
                b.masks.append(mask)
            else:
                issues.append(
                    Issue(
                        "ключевые слова",
                        f"правило заказчика для {b.name} «…{token}» не совпало ни с одним известным креатором бренда "
                        "— не используем, проверьте",
                    )
                )
        for mask in derive_masks(creators, exclude=shop_segments - set(rule)):
            # маска из данных шире правила заказчика с тем же сегментом — правило убираем
            token = mask.strip("*").replace("[._]", "")
            b.masks = [m for m in b.masks if m != customer_mask(token)]
            if mask not in b.masks:
                b.masks.append(mask)
        for url in urls:
            host = urlparse(url).hostname or ""
            if host and key(b.name) and key(b.name) in key(host):
                b.domains.append(host.removeprefix("www."))
        b.domains = sorted(set(b.domains))


def plan_summary(plan: ImportPlan) -> str:
    lines = ["# Импорт списков заказчика", ""]
    lines += [
        "| Бренд | Наш | Аккаунт | Креаторов | Маски ников | Сайт | WW / WB |",
        "|---|---|---|---|---|---|---|",
    ]
    for b in plan.brands:
        lines.append(
            f"| {b.name} | {'да' if b.is_own else ''} | {'@' + b.shop.username if b.shop else '—'} | {len(b.creators)} | "
            f"{', '.join(f'`{m}`' for m in b.masks) or '—'} | {', '.join(b.domains) or '—'} | "
            f"{len(b.ww_codes)} / {len(b.wb_articles)} |"
        )
    warn = [i for i in plan.issues if i.level == "warn"]
    info = [i for i in plan.issues if i.level != "warn"]
    lines += ["", f"## Нужно проверить ({len(warn)})", ""]
    lines += [f"- **{i.where}** — {i.message}" for i in warn] or ["- нет"]
    lines += ["", f"## К сведению ({len(info)})", ""]
    lines += [f"- **{i.where}** — {i.message}" for i in info] or ["- нет"]
    if plan.skipped_sheets:
        lines += [
            "",
            "## Не используем",
            "",
            f"Листы без данных Instagram: {', '.join(plan.skipped_sheets)}.",
        ]
    return "\n".join(lines) + "\n"


# --- запись в базу --------------------------------------------------------------------


async def apply_plan(session: Any, plan: ImportPlan) -> dict[str, int]:
    """Бренды, артикулы и связи. Известный аккаунт связывается сразу, новый — через задачу profile."""
    from sqlalchemy import func, select, text
    from sqlalchemy.dialects.postgresql import insert

    from core.db.models import Account, Brand, BrandAccount, BrandArticle
    from parser.queue import tasks as q

    stats = Counter[str]()
    for bp in plan.brands:
        brand = await session.scalar(select(Brand).where(func.lower(Brand.name) == bp.name.lower()))
        if brand is None:
            brand = Brand(
                name=bp.name, is_own=bp.is_own, hashtags=[], username_masks=[], domains=[]
            )
            session.add(brand)
            await session.flush()
            stats["brands_created"] += 1
        brand.is_own = bp.is_own
        brand.username_masks = sorted(set(brand.username_masks or []) | set(bp.masks))
        brand.domains = sorted(set(brand.domains or []) | set(bp.domains))
        rows = [
            {"brand_id": brand.id, "kind": "ww", "value": v, "source": "xlsx"} for v in bp.ww_codes
        ]
        rows += [
            {"brand_id": brand.id, "kind": "wb", "value": v, "source": "xlsx"}
            for v in bp.wb_articles
        ]
        for i in range(0, len(rows), 1000):
            await session.execute(
                insert(BrandArticle).values(rows[i : i + 1000]).on_conflict_do_nothing()
            )
        stats["articles"] += len(rows)

        refs = ([bp.shop] if bp.shop else []) + bp.creators
        for ref in refs:
            link = {
                "brand_id": brand.id,
                "role": ref.role,
                "is_own_creator": bp.is_own and ref.role == "creator",
                "person_name": ref.person_name,
                "curator": ref.curator,
                "link_type": ref.link_type,
            }
            acc = await session.scalar(select(Account).where(Account.username == ref.username))
            if acc is not None:
                await session.execute(
                    insert(BrandAccount)
                    .values(
                        brand_id=brand.id,
                        user_pk=acc.user_pk,
                        role=ref.role,
                        status="tracked",
                        is_own_creator=link["is_own_creator"],
                        person_name=ref.person_name,
                        curator=ref.curator,
                        link_type=ref.link_type,
                        confidence=1,
                        evidence=[],
                        all_posts_count_for_brand=ref.role == "creator",
                        first_seen=text("current_date"),
                    )
                    .on_conflict_do_update(
                        index_elements=["brand_id", "user_pk"],
                        set_={
                            "status": "tracked",
                            "role": ref.role,
                            "person_name": ref.person_name,
                            "curator": ref.curator,
                        },
                    )
                )
                if ref.role == "shop":
                    brand.shop_user_pk = acc.user_pk
                stats["linked"] += 1
            else:
                task_id = await q.enqueue(
                    session,
                    q.TaskType.PROFILE,
                    username=ref.username,
                    window=f"import-b{brand.id}",
                    priority=q.Priority.P1 if ref.role == "shop" or bp.is_own else q.Priority.P2,
                    payload={"link": link},
                )
                stats["queued" if task_id else "already_queued"] += 1
    await session.commit()
    return dict(stats)
